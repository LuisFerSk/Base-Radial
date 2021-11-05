import errno
from os import mkdir, path
from random import uniform
from math import exp, log, sqrt
from openpyxl import load_workbook
from numpy import random, abs, array, concatenate
from pandas import read_excel, DataFrame, ExcelWriter


def GenerateRadialBase(min, max, row, col):
    return random.uniform(min, max, [row, col])


def EuclideanDistance(entradas, matrizBasesRadiales):
    distanciasEuclidianas = []
    for basesRadiales in matrizBasesRadiales:
        sumatoria = []
        for entrada, baseRadial in zip(entradas, basesRadiales):
            sumatoria.append(pow((entrada - baseRadial), 2))
        distanciasEuclidianas.append(pow(sum(sumatoria), 0.5))
    return distanciasEuclidianas


def RadialBaseFunction(matrizDistanciasEuclidianas):
    funcionesActivacion = []
    for distanciasEuclidianas in matrizDistanciasEuclidianas:
        funcionActivacion = []
        for distanciaEuclidiana in distanciasEuclidianas:
            funcionActivacion.append(
                pow(distanciaEuclidiana, 2) * log(distanciaEuclidiana))
        funcionesActivacion.append(funcionActivacion)
    return funcionesActivacion


def GaussianFunction(matrizDistanciasEuclidianas):
    funcionesActivacion = []
    for distanciasEuclidianas in matrizDistanciasEuclidianas:
        funcionActivacion = []
        for distanciaEuclidiana in distanciasEuclidianas:
            funcionActivacion.append(exp(-pow(distanciaEuclidiana, 2)))
        funcionesActivacion.append(funcionActivacion)
    return funcionesActivacion


def MultiquadraticFunction(matrizDistanciasEuclidianas):
    funcionesActivacion = []
    for distanciasEuclidianas in matrizDistanciasEuclidianas:
        funcionActivacion = []
        for distanciaEuclidiana in distanciasEuclidianas:
            funcionActivacion.append(sqrt(1 + pow(distanciaEuclidiana, 2)))
        funcionesActivacion.append(funcionActivacion)
    return funcionesActivacion


def InverseMultiquadraticFunction(matrizDistanciasEuclidianas):
    funcionesActivacion = []
    for distanciasEuclidianas in matrizDistanciasEuclidianas:
        funcionActivacion = []
        for distanciaEuclidiana in distanciasEuclidianas:
            funcionActivacion.append(
                1 / sqrt(1 + pow(distanciaEuclidiana, 2)))
        funcionesActivacion.append(funcionActivacion)
    return funcionesActivacion


def CalculateOutput(funcionesActivacion, interp):
    salida = []
    for funcionActivacion in funcionesActivacion:
        sumatoria = []
        for fa, ip in zip(funcionActivacion, interp):
            sumatoria.append(fa*ip[0])
        salida.append(sum(sumatoria))
    return salida


def ActivationFunction(funcionActivacion, distanciasEuclidianas):
    switcher = {
        'BASERADIAL': RadialBaseFunction(distanciasEuclidianas),
        'GAUSSIANA': GaussianFunction(distanciasEuclidianas),
        'MULTICUADRATICA': MultiquadraticFunction(distanciasEuclidianas),
        'MC_INVERSA': InverseMultiquadraticFunction(distanciasEuclidianas),
    }
    return switcher.get(funcionActivacion, "ERROR")


def LinearError(salidas, _salida):
    error = []
    entrenamiento = []
    for salida, _salida in zip(salidas, _salida):
        entrenamiento.append([salida[0], _salida])
        error.append(salida[0] - _salida)
    return (error, entrenamiento)


def ErrorG(errorLineal):
    error = 0
    for salida in errorLineal:
        error += abs(salida)
    return error / len(errorLineal)


def ReadData(ruta):
    entradas = []
    salidas = []

    ejercicio = path.basename(path.splitext(ruta)[0])

    workbook = load_workbook(ruta)

    matriz = read_excel(ruta, sheet_name='Matriz')
    for i in range(len(matriz.columns)):
        entradas.append([fila[i] for fila in matriz.to_numpy(
        )]) if 'X' in matriz.columns[i] else salidas.append([fila[i] for fila in matriz.to_numpy()])

    matrizBaseRadiales = []
    funcionActivacion = 'BASERADIAL'
    neuronas = int(uniform(1, 9))
    error = 0.001

    if 'Config' in workbook.sheetnames:
        matrizBaseRadiales = read_excel(
            ruta, sheet_name='Bases Radiales').to_numpy()
        funcionActivacion = read_excel(
            ruta, sheet_name='Config').to_numpy()[0][0]
        neuronas = read_excel(
            ruta, sheet_name='Config').to_numpy()[0][2]
        error = read_excel(ruta, sheet_name='Config').to_numpy()[0][1]

    return (ejercicio, matriz, array(entradas).transpose(), array(salidas).transpose(), matrizBaseRadiales, funcionActivacion, neuronas, error)


def SaveResults(ejercicio, entrenamiento, entradas, salidas, basesRadiales, funcionSalida, error, neuronas):
    dfMatrix = DataFrame(concatenate((array(entradas), array(salidas)), axis=1), columns=[
        'X' + str(x+1) for x in range(len(entradas[0]))] + ['YD' + str(x+1) for x in range(len(salidas[0]))])
    dfBasesRadiales = DataFrame(basesRadiales, columns=[
        'BR' + str(x+1) for x in range(len(basesRadiales[0]))])
    dfConfig = DataFrame([[funcionSalida, error, neuronas]], columns=[
        'Func Activacion', 'Error Maximo', 'Neuronas'])

    try:
        mkdir('src/DATA/' + entrenamiento)
    except OSError as e:
        if e.errno != errno.EEXIST:
            raise

    try:
        mkdir('src/DATA/' + entrenamiento + '/' + funcionSalida)
    except OSError as e:
        if e.errno != errno.EEXIST:
            raise

    raiz = 'src/DATA/' + entrenamiento + '/' + funcionSalida + '/' + ejercicio + '.xlsx' if 'out' == entrenamiento else 'src/DATA/' + \
        entrenamiento + '/' + funcionSalida + '/' + ejercicio + ' - ' + error + '.xlsx'

    with ExcelWriter(raiz) as writer:  # pylint: disable=abstract-class-instantiated
        dfMatrix.to_excel(writer, sheet_name='Matriz', index=False)
        dfBasesRadiales.to_excel(
            writer, sheet_name='Bases Radiales', index=False)
        dfConfig.to_excel(writer, sheet_name='Config', index=False)
